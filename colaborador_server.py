#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Servidor local para a aba "Colaborador" do painel: concluir atividades,
justificar pendências e cadastrar novas tarefas direto pelo navegador,
escrevendo na hora na planilha ProgramaçãoSMA.xlsx e devolvendo os dados
atualizados para o painel se redesenhar sozinho — sem precisar rodar o
atualizar_painel.bat a cada clique.

Uso:
    python colaborador_server.py
    (ou dê duplo clique em abrir_colaborador.bat)

Abre http://localhost:5678 sozinho. Só funciona enquanto esta janela
estiver aberta — feche-a (ou Ctrl+C) para encerrar.

Não tem login: qualquer pessoa com acesso a este endereço local pode
editar qualquer atividade. O campo "Meu nome" na aba Colaborador é só
para filtrar a visualização, não é uma trava de acesso — ver README.

Publicar as mudanças no link público (GitHub Pages) continua sendo um
passo separado e deliberado: rode atualizar_painel.bat quando quiser
levar a versão mais recente da planilha para lá.
"""
import datetime
import json
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

try:
    import openpyxl
except ImportError:
    sys.exit("Faltou instalar a dependência: pip install -r requirements.txt")

from build_dashboard import achar_coluna, carregar_dados, montar_html
from update_and_publish import XLSX_PATH

PORT = 5678
ABA_ATIVIDADES = "Atividades(SM&A)"
ABA_CARGA = "Carga Diária"

# Todo acesso de escrita à planilha passa por este lock: como o servidor
# atende cada requisição numa thread própria, sem isto duas ações quase
# simultâneas (dois cliques em "Concluir" ao mesmo tempo, por exemplo)
# poderiam ler/escrever o arquivo ao mesmo tempo e corrompê-lo.
_lock = threading.Lock()

CANDIDATOS_COLUNA = {
    "numero": ["N°", "Nº", "N", "NUMERO", "NÚMERO"],
    "desc": ["DESCRIÇÃO DAS ATIVIDADES", "DESCRIÇÃO"],
    "local": ["LOCAL"],
    "fim": ["DATA PREVISTA CONCLUSÃO", "DATA PREVISTA CONCLUSAO"],
    "horas": ["PREVISÃO DE HORAS", "PREVISAO DE HORAS"],
    "situacao": ["SITUAÇÃO DO PRAZO", "SITUACAO DO PRAZO"],
    "status": ["STATUS"],
    "resp": ["RESPONSÁVEL", "RESPONSAVEL", "COLABORADOR"],
    "justificativa": ["JUSTIFICATIVA", "MOTIVO", "MOTIVO DO ATRASO"],
}


def _mapear_colunas(headers):
    return {chave: achar_coluna(headers, candidatos) for chave, candidatos in CANDIDATOS_COLUNA.items()}


def _garantir_coluna(ws, headers, chave, titulo_padrao):
    """Índice de coluna (1-based, para uso direto em ws.cell) para
    `chave`; se a planilha ainda não tem essa coluna, cria uma nova no
    fim com o cabeçalho `titulo_padrao` — assim RESPONSÁVEL/JUSTIFICATIVA
    passam a existir sozinhas na primeira vez que alguém as usa, sem
    precisar editar a planilha à mão antes."""
    idx0 = achar_coluna(headers, CANDIDATOS_COLUNA[chave])
    if idx0 is not None:
        return idx0 + 1
    nova_col = len(headers) + 1
    ws.cell(row=1, column=nova_col, value=titulo_padrao)
    headers.append(titulo_padrao)
    return nova_col


def _abrir_planilha():
    # Sem data_only: preserva fórmulas que já existam em outras células
    # (só mexemos nas células específicas que cada ação altera).
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb[ABA_ATIVIDADES]
    headers = [c.value for c in ws[1]]
    return wb, ws, headers


def _linha_por_numero(ws, col0_numero, numero):
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col0_numero + 1).value
        if v is not None and str(v).strip() == str(numero).strip():
            return r
    return None


def concluir_atividade(numero):
    with _lock:
        wb, ws, headers = _abrir_planilha()
        cols = _mapear_colunas(headers)
        if cols["numero"] is None or cols["status"] is None:
            raise RuntimeError("Não encontrei as colunas N° / STATUS na planilha.")
        linha = _linha_por_numero(ws, cols["numero"], numero)
        if linha is None:
            raise RuntimeError(f"Atividade N° {numero} não encontrada na planilha.")

        # O resto do painel compara STATUS com a string exata "concluído"
        # (minúsculo) — ver template.html — por isso a grafia aqui não é
        # cosmética, é o valor que o JS de fato reconhece.
        ws.cell(row=linha, column=cols["status"] + 1, value="concluído")
        if cols["situacao"] is not None:
            cel_situacao = ws.cell(row=linha, column=cols["situacao"] + 1)
            if cel_situacao.data_type != "f":  # não sobrescreve fórmula
                cel_situacao.value = "Concluído"

        wb.save(XLSX_PATH)
    return carregar_dados(XLSX_PATH, ABA_ATIVIDADES, ABA_CARGA)


def justificar_atividade(numero, texto):
    with _lock:
        wb, ws, headers = _abrir_planilha()
        cols = _mapear_colunas(headers)
        if cols["numero"] is None:
            raise RuntimeError("Não encontrei a coluna N° na planilha.")
        linha = _linha_por_numero(ws, cols["numero"], numero)
        if linha is None:
            raise RuntimeError(f"Atividade N° {numero} não encontrada na planilha.")

        col_just = _garantir_coluna(ws, headers, "justificativa", "JUSTIFICATIVA")
        ws.cell(row=linha, column=col_just, value=texto)

        wb.save(XLSX_PATH)
    return carregar_dados(XLSX_PATH, ABA_ATIVIDADES, ABA_CARGA)


def nova_atividade(desc, local, fim, horas, resp):
    if not desc:
        raise RuntimeError("Descreva a atividade antes de adicionar.")
    fim_date = datetime.date.fromisoformat(fim) if fim else None
    try:
        horas_num = float(horas) if horas else 0
    except (TypeError, ValueError):
        horas_num = 0

    with _lock:
        wb, ws, headers = _abrir_planilha()
        cols = _mapear_colunas(headers)

        maior_numero = 0
        if cols["numero"] is not None:
            for r in range(2, ws.max_row + 1):
                v = ws.cell(row=r, column=cols["numero"] + 1).value
                try:
                    maior_numero = max(maior_numero, int(v))
                except (TypeError, ValueError):
                    pass
        novo_numero = maior_numero + 1

        col_resp = _garantir_coluna(ws, headers, "resp", "RESPONSÁVEL")
        cols = _mapear_colunas(headers)  # recalcula: _garantir_coluna pode ter criado colunas novas

        linha_nova = ws.max_row + 1
        if cols["numero"] is not None:
            ws.cell(row=linha_nova, column=cols["numero"] + 1, value=novo_numero)
        if cols["desc"] is not None:
            ws.cell(row=linha_nova, column=cols["desc"] + 1, value=desc)
        if cols["local"] is not None and local:
            ws.cell(row=linha_nova, column=cols["local"] + 1, value=local)
        if cols["fim"] is not None and fim_date:
            ws.cell(row=linha_nova, column=cols["fim"] + 1, value=fim_date)
        if cols["horas"] is not None and horas_num:
            ws.cell(row=linha_nova, column=cols["horas"] + 1, value=horas_num)
        if cols["status"] is not None:
            ws.cell(row=linha_nova, column=cols["status"] + 1, value="A programar")
        ws.cell(row=linha_nova, column=col_resp, value=resp)

        wb.save(XLSX_PATH)
    return carregar_dados(XLSX_PATH, ABA_ATIVIDADES, ABA_CARGA)


class ColaboradorHandler(BaseHTTPRequestHandler):
    def _enviar_json(self, obj, status=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _enviar_html(self, html):
        body = html.encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _ler_json(self):
        length = int(self.headers.get("Content-Length", 0))
        raw = self.rfile.read(length) if length else b"{}"
        return json.loads(raw or b"{}")

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            dados = carregar_dados(XLSX_PATH, ABA_ATIVIDADES, ABA_CARGA)
            self._enviar_html(montar_html(dados))
        elif self.path == "/api/dados":
            self._enviar_json(carregar_dados(XLSX_PATH, ABA_ATIVIDADES, ABA_CARGA))
        else:
            self.send_error(404)

    def do_POST(self):
        try:
            payload = self._ler_json()
            if self.path == "/api/concluir":
                dados = concluir_atividade(payload["numero"])
            elif self.path == "/api/justificativa":
                dados = justificar_atividade(payload["numero"], (payload.get("texto") or "").strip())
            elif self.path == "/api/nova-atividade":
                dados = nova_atividade(
                    (payload.get("desc") or "").strip(),
                    (payload.get("local") or "").strip(),
                    payload.get("fim") or None,
                    payload.get("horas"),
                    (payload.get("resp") or "").strip(),
                )
            else:
                self.send_error(404)
                return
            self._enviar_json(dados)
        except Exception as exc:
            self._enviar_json({"erro": str(exc)}, status=400)


def main():
    from pathlib import Path

    print(f"Planilha: {XLSX_PATH}")
    if not Path(XLSX_PATH).exists():
        print("ERRO: não encontrei a planilha nesse caminho.")
        print("Confira/ajuste XLSX_PATH em update_and_publish.py.")
        sys.exit(1)

    server = ThreadingHTTPServer(("127.0.0.1", PORT), ColaboradorHandler)
    url = f"http://localhost:{PORT}/"
    print("=" * 60)
    print("Painel do Colaborador no ar!")
    print("=" * 60)
    print(f"Abrindo {url}")
    print("Deixe esta janela aberta enquanto estiver usando.")
    print("Feche a janela (ou Ctrl+C) para encerrar.\n")
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.shutdown()


if __name__ == "__main__":
    main()
