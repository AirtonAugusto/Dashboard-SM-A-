#!/usr/bin/env python3
"""
Gera o dashboard HTML (Programação de Atividades SM&A) a partir da planilha
ProgramaçãoSMA.xlsx.

Uso:
    python build_dashboard.py --xlsx caminho/para/ProgramacaoSMA.xlsx \
                               --logo assets/logo_anglogold.png \
                               --template template.html \
                               --out docs/index.html

Basta rodar este script de novo, apontando para a planilha atualizada,
sempre que quiser refletir os dados mais recentes no painel. O arquivo
gerado em --out é autocontido (HTML + CSS + JS + dados + logo embutidos).
O padrão é docs/index.html porque é o caminho que o GitHub Pages usa
para publicar o painel automaticamente em um link público (veja o
README) — mas também pode ser aberto direto no navegador, sem precisar
de internet.

As funções carregar_dados()/montar_html() também são reaproveitadas por
colaborador_server.py (aba "Colaborador" do painel, ver README) para
gerar a mesma página ao vivo, sem passar por um arquivo intermediário.
"""
import argparse
import base64
import datetime
import json
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Faltou instalar a dependência: pip install openpyxl")


def normalizar_texto(s):
    """Maiúsculas, sem espaço nas pontas, sem acento — para comparar
    cabeçalhos de coluna sem depender de grafia exata (espaço a mais,
    acento diferente etc.)."""
    if s is None:
        return ""
    s = str(s).strip().upper()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def achar_coluna(headers, candidatos):
    """Índice (0-based) do primeiro cabeçalho em `headers` que bater com
    algum nome em `candidatos`, ignorando acento/maiúscula/espaço extra.
    None se nenhum bater — quem chama decide o que fazer (campo em
    branco, criar a coluna etc.)."""
    normalizados = {normalizar_texto(h): i for i, h in enumerate(headers) if h is not None}
    for cand in candidatos:
        idx = normalizados.get(normalizar_texto(cand))
        if idx is not None:
            return idx
    return None


def extrair_atividades(ws):
    headers = [c.value for c in ws[1]]
    # RESPONSÁVEL e JUSTIFICATIVA são campos novos (aba Colaborador, ver
    # README) — usam busca por nome aproximado porque a planilha pode não
    # ter essas colunas ainda (nesse caso ficam em branco) ou ter o
    # cabeçalho escrito de um jeito ligeiramente diferente.
    idx_resp = achar_coluna(headers, ["RESPONSÁVEL", "RESPONSAVEL", "COLABORADOR"])
    idx_just = achar_coluna(headers, ["JUSTIFICATIVA", "MOTIVO", "MOTIVO DO ATRASO"])

    acts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(c is None for c in row):
            continue
        d = dict(zip(headers, row))
        ini = d.get("DATA PREVISTA INICIO ")
        fim = d.get("DATA PREVISTA CONCLUSÃO")
        resp = row[idx_resp] if idx_resp is not None and idx_resp < len(row) else None
        justificativa = row[idx_just] if idx_just is not None and idx_just < len(row) else None
        acts.append({
            "id": d.get("N°"),
            "desc": (d.get("DESCRIÇÃO DAS ATIVIDADES") or "").strip(),
            "local": d.get("LOCAL"),
            "ini": ini.strftime("%Y-%m-%d") if isinstance(ini, (datetime.date, datetime.datetime)) else None,
            "fim": fim.strftime("%Y-%m-%d") if isinstance(fim, (datetime.date, datetime.datetime)) else None,
            "horas": d.get("PREVISÃO DE HORAS") or 0,
            "situacao": (d.get("SITUAÇÃO DO PRAZO ") or "").strip(),
            "status": (d.get("STATUS") or "").strip(),
            "obs": (d.get("OBSERVAÇÕES") or "").strip() if d.get("OBSERVAÇÕES") else None,
            "resp": (resp or "").strip(),
            "justificativa": (justificativa or "").strip(),
        })
    return acts


def extrair_carga_diaria(ws):
    carga = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or not isinstance(row[0], (datetime.date, datetime.datetime)):
            continue
        carga.append({
            "data": row[0].strftime("%Y-%m-%d"),
            "tarefas": row[1],
            "horas": round(row[2], 2) if row[2] is not None else 0,
            "situacao": row[3],
        })
    return carga


def carregar_dados(xlsx_path, aba_atividades="Atividades(SM&A)", aba_carga="Carga Diária"):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    acts = extrair_atividades(wb[aba_atividades])
    carga = extrair_carga_diaria(wb[aba_carga]) if aba_carga in wb.sheetnames else []
    return {"acts": acts, "carga": carga}


def montar_html(dados, logo_path="assets/logo_anglogold.png", template_path="template.html"):
    data_json = json.dumps(dados, ensure_ascii=False, separators=(",", ":"))
    data_json = data_json.replace("</script>", "<\\/script>")

    logo_b64 = base64.b64encode(Path(logo_path).read_bytes()).decode() if Path(logo_path).exists() else ""

    template = Path(template_path).read_text(encoding="utf-8")
    return template.replace("__DATA_JSON__", data_json).replace("__LION_B64__", logo_b64)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="Caminho da planilha ProgramaçãoSMA.xlsx")
    ap.add_argument("--logo", default="assets/logo_anglogold.png", help="Logo (PNG) para a barra lateral")
    ap.add_argument("--template", default="template.html", help="Template HTML com os placeholders")
    ap.add_argument("--out", default="docs/index.html", help="Arquivo HTML final a gerar")
    ap.add_argument("--aba-atividades", default="Atividades(SM&A)", help="Nome da aba de atividades")
    ap.add_argument("--aba-carga", default="Carga Diária", help="Nome da aba de carga diária")
    args = ap.parse_args()

    dados = carregar_dados(args.xlsx, args.aba_atividades, args.aba_carga)
    out_html = montar_html(dados, args.logo, args.template)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(out_html, encoding="utf-8")

    print(f"OK: {len(dados['acts'])} atividades, {len(dados['carga'])} dias de carga diária.")
    print(f"Painel gerado em: {out_path.resolve()}")


if __name__ == "__main__":
    main()
